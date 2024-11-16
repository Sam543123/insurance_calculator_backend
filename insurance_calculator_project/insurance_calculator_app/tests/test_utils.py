import openpyxl


def compare_excel_files(file_like_obj1, file_like_obj2):
    """Check if two excel files are equal"""
    workbook1 = openpyxl.load_workbook(file_like_obj1, data_only=True)
    work_sheet1 = workbook1.active
    workbook2 = openpyxl.load_workbook(file_like_obj2, data_only=True)
    work_sheet2 = workbook2.active
    if work_sheet1.max_row != work_sheet2.max_row or work_sheet1.max_column != work_sheet2.max_column:
        return False
    for row1, row2 in zip(work_sheet1.iter_rows(values_only=True), work_sheet2.iter_rows(values_only=True)):
        if row1 != row2:
            return False
    return True
